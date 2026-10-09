"""The 채점결과.xlsx workbook: results sheet, analysis sheets and the colour legend."""

from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from quiz_reporter.errors import Ok
from quiz_reporter.quiz.grading import bank_from_sheet, sheet_from_scores, suggest_cutoff
from quiz_reporter.quiz.pipeline import Graded, grade
from quiz_reporter.quiz.report import build_reports
from quiz_reporter.quiz.responses import FormResponses, read_form_responses
from quiz_reporter.quiz.result_book import RESULT_FILENAME, SHEET_NAMES, result_workbook_bytes
from tests.helpers.quiz_forms import ROWS, form_csv, full_bank

PINK = "FFFDE2E4"
GREY = "FFE7E6E6"
YELLOW = "FFFFEB9C"
HEADER_ROW = 3
FIRST_STUDENT_ROW = 4


def _responses(tmp_path: Path, rows=ROWS) -> FormResponses:
    path = tmp_path / "응답.csv"
    path.write_bytes(form_csv(rows))
    responses = read_form_responses(str(path))
    assert isinstance(responses, Ok), responses
    return responses.value


def _graded(tmp_path: Path, rows=ROWS, *, cutoff: bool = True, bank=None) -> Graded:
    responses = _responses(tmp_path, rows)
    result = grade(responses, suggest_cutoff(responses) if cutoff else None, bank or full_bank())
    assert isinstance(result, Ok), result
    return result.value


def _book(graded: Graded):
    return load_workbook(BytesIO(result_workbook_bytes(graded, "생리 퀴즈")))


@pytest.fixture
def graded(tmp_path) -> Graded:
    return _graded(tmp_path)


@pytest.fixture
def sheet(graded):
    return _book(graded)["채점결과"]


def _rgb(cell) -> str | None:
    return cell.fill.start_color.rgb if cell.fill.fill_type == "solid" else None


def test_the_filename_and_sheet_order_with_excluded_rows(graded):
    assert RESULT_FILENAME == "채점결과.xlsx"
    assert graded.excluded
    assert _book(graded).sheetnames == list(SHEET_NAMES)


def test_the_excluded_sheet_is_left_out_when_nothing_was_excluded(tmp_path):
    graded = _graded(tmp_path, ROWS[:4], cutoff=False)
    assert not graded.excluded

    names = _book(graded).sheetnames

    assert names == [name for name in SHEET_NAMES if name != "제외된 응답"]


def test_the_title_note_and_header_rows(sheet):
    assert sheet["A1"].value == "생리 퀴즈 · 채점결과"
    assert sheet["A1"].font.bold and sheet["A1"].font.size > 11
    assert (
        sheet["A2"].value
        == "이름 가나다순 · 문항 칸은 고른 보기 번호(1~5) · 분홍 = 틀린 답 · 회색 = 답 없음"
        " · 노랑 = 학번 확인 필요 · 정답이 빈 문항은 채점하지 않음"
    )
    header = [cell.value for cell in sheet[HEADER_ROW]]
    assert header == ["순번", "학번", "이름", "점수", "만점", 1, 2, 3, "비고"]
    assert sheet.freeze_panes == "D4"


def test_one_row_per_student_in_serial_order_with_the_report_scores(graded, sheet):
    rows = [
        [cell.value for cell in sheet[line][:5]]
        for line in range(FIRST_STUDENT_ROW, FIRST_STUDENT_ROW + 4)
    ]

    assert [row[2] for row in rows] == ["가나", "다라", "마바", "사아"]
    assert [row[0] for row in rows] == [1, 2, 3, 4]
    assert [row[3] for row in rows] == [report.score for report in graded.reports] == [3, 1, 0, 2]
    assert {row[4] for row in rows} == {3}


def test_ids_and_names_are_text(sheet):
    cell = sheet.cell(FIRST_STUDENT_ROW, 2)

    assert cell.value == "20260001"
    assert cell.data_type == "s"
    assert sheet.cell(FIRST_STUDENT_ROW, 3).data_type == "s"


def test_answers_show_the_chosen_option_and_wrong_ones_are_pink(sheet):
    right = [sheet.cell(FIRST_STUDENT_ROW, column) for column in (6, 7, 8)]
    assert [cell.value for cell in right] == [3, 2, 3]
    assert [_rgb(cell) for cell in right] == [None, None, None]
    assert all(cell.alignment.horizontal == "center" for cell in right)

    mixed = [sheet.cell(FIRST_STUDENT_ROW + 1, column) for column in (6, 7, 8)]
    assert [cell.value for cell in mixed] == [2, 3, 3]
    assert [_rgb(cell) for cell in mixed] == [PINK, PINK, None]


def test_a_student_id_that_is_not_eight_digits_is_yellow_with_a_note(sheet):
    line = FIRST_STUDENT_ROW + 3
    assert sheet.cell(line, 3).value == "사아"

    assert not sheet.cell(line, 2).value
    assert _rgb(sheet.cell(line, 2)) == YELLOW
    assert "2026004" in sheet.cell(line, 9).value
    assert "8자리 숫자 아님" in sheet.cell(line, 9).value
    assert sheet.cell(FIRST_STUDENT_ROW, 9).value is None


def test_a_missing_id_says_so(graded):
    students = (replace(graded.students[0], student_id="", typed_id=""), *graded.students[1:])
    reports, summary = build_reports(graded.bank, students)

    sheet = _book(replace(graded, students=students, reports=reports, summary=summary))["채점결과"]

    assert sheet.cell(FIRST_STUDENT_ROW, 9).value == "학번 없음"


def test_the_answer_and_rate_rows_follow_a_blank_row(graded, sheet):
    last = FIRST_STUDENT_ROW + 3
    assert all(cell.value is None for cell in sheet[last + 1])
    answer_row, rate_row = last + 2, last + 3

    assert sheet.cell(answer_row, 3).value == "정답"
    assert [sheet.cell(answer_row, column).value for column in (6, 7, 8)] == [3, 2, 3]
    assert sheet.cell(rate_row, 3).value == "정답률"
    rates = [sheet.cell(rate_row, column) for column in (6, 7, 8)]
    assert [cell.value for cell in rates] == list(graded.summary.correct_rate)
    assert {cell.number_format for cell in rates} == {"0%"}


def test_a_blank_answer_is_grey_and_counts_as_missed(graded):
    first = graded.students[2]  # 마바
    students = (
        *graded.students[:2],
        replace(first, choices=(None, *first.choices[1:])),
        *graded.students[3:],
    )
    reports, summary = build_reports(graded.bank, students)

    sheet = _book(replace(graded, students=students, reports=reports, summary=summary))["채점결과"]
    cell = sheet.cell(FIRST_STUDENT_ROW + 2, 6)

    assert cell.value is None
    assert _rgb(cell) == GREY
    assert _rgb(sheet.cell(FIRST_STUDENT_ROW + 2, 7)) == PINK


def test_a_question_without_a_known_answer_is_not_coloured_or_graded(tmp_path):
    rows = (
        (0, "20260001", "가나", (3, 2, 2)),
        (1, "20260002", "다라", (2, 3, 2)),
        (2, "20260003", "마바", (1, 1, 1)),
        (3, "20260004", "사아", (3, 2, 1)),
    )
    responses = _responses(tmp_path, rows)
    answer_sheet = sheet_from_scores(responses).value
    # Options are numbered in the order students chose them, so the answers are 1, 1.
    assert answer_sheet.answers == (1, 1, None)
    bank = bank_from_sheet(responses, answer_sheet)
    graded = _graded(tmp_path, rows, cutoff=False, bank=bank)

    sheet = _book(graded)["채점결과"]
    # A bank from CSV scores numbers options by the responses, so a note row comes first.
    assert bank.response_order
    students = range(5, 9)

    assert [sheet.cell(line, 5).value for line in students] == [2, 2, 2, 2]
    assert [sheet.cell(line, 4).value for line in students] == [2, 0, 0, 2]
    shown = [sheet.cell(line, 8).value for line in students]
    assert shown == [1, 1, 2, 2]
    assert [_rgb(sheet.cell(line, 8)) for line in students] == [None] * 4
    assert sheet.cell(10, 8).value is None
    assert sheet.cell(11, 8).value is None
    assert sheet.cell(11, 6).value == graded.summary.correct_rate[0]


def test_the_analysis_sheet_keeps_the_trap_analysis(graded):
    analysis = _book(graded)["문항 분석"]

    header = [cell.value for cell in analysis[1]]
    assert header[:5] == ["번호", "단원", "문제", "정답", "정답률"]
    assert header[-2:] == ["가장 많이 고른 오답", "그 함정 유형"]
    first = [cell.value for cell in analysis[2]]
    assert first[0] == 1
    assert first[3] == 3
    assert first[4] == 0.5
    assert first[5:10] == [1, 1, 2, 0, 0]
    assert _rgb(analysis.cell(2, 8)) == "FFC6EFCE"  # the right option's count


def test_the_other_sheets_are_filled(graded):
    book = _book(graded)

    trap_row = [cell.value for cell in book["학생별 함정"][2]]
    assert trap_row[:5] == [1, "20260001", "가나", 3, 3]
    assert book["학생별 함정"].cell(2, 2).data_type == "s"
    assert [cell.value for cell in book["반 전체 함정"][1]] == ["함정 유형", "선택 횟수(오답)"]
    assert [cell.value for cell in book["제외된 응답"][1]] == [
        "시각",
        "학번(입력값)",
        "이름",
        "이유",
    ]
    assert book["제외된 응답"].max_row == 1 + len(graded.excluded)
    assert book["문항표"].max_row > 1


def test_the_legend_lists_each_colour_used(graded):
    legend = _book(graded)["색 설명"]

    swatches = [_rgb(legend.cell(line, 1)) for line in range(2, 7)]
    assert swatches == [PINK, GREY, YELLOW, "FFC6EFCE", "FFFFC7CE"]
    assert legend.cell(2, 2).value.startswith("분홍")


def test_every_header_fits_its_column(graded):
    """No header is cut off in Excel: each column is at least as wide as its header text."""
    from openpyxl.utils import get_column_letter

    from quiz_reporter.quiz.result_book import text_width

    workbook = _book(graded)
    header_rows = {"채점결과": 3}
    for sheet in workbook.worksheets:
        if sheet.title == "색 설명":
            continue
        row = header_rows.get(sheet.title, 1)
        for cell in sheet[row]:
            if cell.value is None:
                continue
            width = sheet.column_dimensions[get_column_letter(cell.column)].width
            assert width >= text_width(cell.value), (sheet.title, cell.value, width)


def test_a_quiz_name_that_looks_like_a_formula_stays_text(graded):
    book = load_workbook(BytesIO(result_workbook_bytes(graded, "=1단원 확인")))

    cell = book["채점결과"]["A1"]
    assert cell.data_type == "s"
    assert cell.value == "=1단원 확인 · 채점결과"


def test_characters_excel_rejects_are_dropped_from_names(graded):
    from dataclasses import replace

    students = (replace(graded.students[0], name="가\x0b나"), *graded.students[1:])
    reports = (replace(graded.reports[0], student=students[0]), *graded.reports[1:])

    book = _book(replace(graded, students=students, reports=reports))

    assert book["채점결과"]["C4"].value == "가나"


def test_tied_wrong_options_are_all_named():
    from quiz_reporter.quiz.result_book import _most_chosen

    assert _most_chosen([2], 3) == "2번 (3명)"
    assert _most_chosen([3, 5], 2) == "3·5번 (각 2명)"
    assert _most_chosen([], 0) == ""


def test_answers_are_plain_numbers_for_later_processing(sheet):
    """Option numbers are numeric cells (1~5), never circled characters."""
    for row in sheet.iter_rows(min_row=4, max_row=7, min_col=6, max_col=8):
        for cell in row:
            assert cell.value is None or (isinstance(cell.value, int) and 1 <= cell.value <= 5)
            assert cell.data_type == "n"


def test_a_response_order_quiz_says_so_above_both_headers(tmp_path):
    from quiz_reporter.quiz.grading import bank_from_sheet, sheet_from_scores
    from quiz_reporter.quiz.result_book import ORDER_NOTE

    responses = _responses(tmp_path)
    bank = bank_from_sheet(responses, sheet_from_scores(responses).value)
    book = _book(_graded(tmp_path, bank=bank))

    result, analysis = book["채점결과"], book["문항 분석"]
    assert result["A3"].value == ORDER_NOTE
    assert result["A4"].value == "순번"
    assert result.freeze_panes == "D5"
    assert analysis["A1"].value == ORDER_NOTE
    assert analysis["A2"].value == "번호"
    assert analysis.freeze_panes == "D3"
