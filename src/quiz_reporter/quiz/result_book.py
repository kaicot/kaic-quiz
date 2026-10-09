"""The teacher's result workbook: every student's answers and score, then the trap analysis."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from quiz_reporter.quiz.bank import TRAP_NAMES, QuizBank, bank_workbook_bytes
from quiz_reporter.quiz.pipeline import Graded

RESULT_FILENAME = "채점결과.xlsx"
RESULT_SHEET = "채점결과"
ANALYSIS_SHEET = "문항 분석"
SHEET_NAMES = (
    RESULT_SHEET,
    ANALYSIS_SHEET,
    "학생별 함정",
    "반 전체 함정",
    "제외된 응답",
    "문항표",
    "색 설명",
)

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)
_NOTE = Font(color="FF6B7280")
_CENTER = Alignment(horizontal="center", vertical="center")


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


_HEAD = _fill("FFD9E1F2")
_RIGHT = _fill("FFC6EFCE")  # 문항 분석: the right option
_HOT = _fill("FFFFC7CE")  # 문항 분석: a wrong option chosen more than the right one
_WRONG = _fill("FFFDE2E4")  # 채점결과: a wrong answer
_BLANK = _fill("FFE7E6E6")  # 채점결과: left blank
_CHECK = _fill("FFFFEB9C")  # 채점결과: student ID needs checking


def _text(cell: Any, value: str) -> None:
    cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell.data_type = "s"


def _clean(value: object) -> object:
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _no_formulas(workbook: Any) -> None:
    """Quiz names, units and traps are typed by people: '=…' stays text, not a formula."""
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"


def _header(sheet: Any, values: list[object]) -> None:
    """Append ``values`` as the next row and style it as a header."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        cell.font = _BOLD
        cell.fill = _HEAD


def text_width(value: object) -> int:
    """Excel column units a value needs: wide (Hangul, ①) characters count double."""
    text = "" if value is None else str(value)
    return sum(2 if ord(char) >= 0x1100 else 1 for char in text)


def _fit_columns(
    sheet: Any, *, first_row: int = 1, fixed: dict[str, float] | None = None, cap: int = 40
) -> None:
    """Widen each column to its longest text (from ``first_row``), within ``cap``."""
    fixed = fixed or {}
    for column in sheet.iter_cols(min_row=first_row):
        letter = get_column_letter(column[0].column)
        if letter in fixed:
            sheet.column_dimensions[letter].width = fixed[letter]
            continue
        widest = max((text_width(cell.value) for cell in column), default=0)
        sheet.column_dimensions[letter].width = min(cap, max(6, widest + 3))


def _result_sheet(sheet: Any, graded: Graded, title: str) -> None:
    sheet.title = RESULT_SHEET
    items = graded.bank.items
    sheet.append([f"{title} · 채점결과"])
    sheet["A1"].font = _TITLE
    sheet.append(
        [
            "이름 가나다순 · 문항 칸은 고른 보기 번호(1~5) · 분홍 = 틀린 답 · 회색 = 답 없음"
            " · 노랑 = 학번 확인 필요 · 정답이 빈 문항은 채점하지 않음"
        ]
    )
    sheet["A2"].font = _NOTE
    _header(
        sheet, ["순번", "학번", "이름", "점수", "만점", *(item.number for item in items), "비고"]
    )
    header_row = sheet.max_row
    for cell in sheet[header_row]:
        cell.alignment = _CENTER
    first_question = 6
    note_column = first_question + len(items)
    for student, report in zip(graded.students, graded.reports, strict=True):
        sheet.append([student.serial, None, None, report.score, report.maximum])
        line = sheet.max_row
        _text(sheet.cell(line, 2), student.student_id)
        _text(sheet.cell(line, 3), student.name)
        for item in items:
            choice = student.choices[item.number - 1]
            cell = sheet.cell(line, first_question + item.number - 1)
            cell.value = choice
            cell.alignment = _CENTER
            if item.answer is None:
                continue
            if choice is None:
                cell.fill = _BLANK
            elif choice != item.answer:
                cell.fill = _WRONG
        if not student.student_id:
            sheet.cell(line, 2).fill = _CHECK
            sheet.cell(line, note_column).value = (
                f"학번 입력값 '{student.typed_id}' (8자리 숫자 아님)"
                if student.typed_id
                else "학번 없음"
            )
    sheet.append([])
    sheet.append([None, None, "정답"])
    answer_row = sheet.max_row
    sheet.append([None, None, "정답률"])
    rate_row = sheet.max_row
    for item in items:
        column = first_question + item.number - 1
        answer = sheet.cell(answer_row, column)
        # No known answer: left empty so the column stays numbers only.
        answer.value = item.answer
        answer.alignment = _CENTER
        if item.answer is not None:
            rate = sheet.cell(rate_row, column)
            rate.value = graded.summary.correct_rate[item.number - 1]
            rate.number_format = "0%"
            rate.alignment = _CENTER
    for line in (answer_row, rate_row):
        sheet.cell(line, 3).font = _BOLD
    _fit_columns(sheet, first_row=header_row, cap=48)
    sheet.freeze_panes = sheet.cell(header_row + 1, 4)


def _most_chosen(numbers: list[int], count: int) -> str:
    """'2번 (3명)', or every tied option: '3·5번 (각 2명)'."""
    if not numbers:
        return ""
    options = "·".join(str(number) for number in numbers)
    return f"{options}번 ({'각 ' if len(numbers) > 1 else ''}{count}명)"


def _analysis_sheet(sheet: Any, graded: Graded) -> None:
    sheet.title = ANALYSIS_SHEET
    _header(
        sheet,
        [
            "번호",
            "단원",
            "문제",
            "정답",
            "정답률",
            *(f"보기{number} 선택" for number in range(1, 6)),
            "가장 많이 고른 오답",
            "그 함정 유형",
        ],
    )
    for item in graded.bank.items:
        index = item.number - 1
        counts = graded.summary.option_counts[index]
        wrong = [(count, number) for number, count in enumerate(counts, 1) if number != item.answer]
        top_count = max((count for count, _ in wrong), default=0)
        tops = [number for count, number in wrong if top_count and count == top_count]
        top = tops[0] if tops else 0
        sheet.append(
            [
                item.number,
                item.unit,
                item.question,
                item.answer,
                round(graded.summary.correct_rate[index], 3) if item.answer else None,
                *counts,
                _most_chosen(tops, top_count),
                " · ".join(dict.fromkeys(item.traps[n - 1] for n in tops if item.traps[n - 1])),
            ]
        )
        line = sheet.max_row
        _text(sheet.cell(line, 3), item.question)
        sheet.cell(line, 5).number_format = "0%"
        if item.answer:
            sheet.cell(line, 5 + item.answer).fill = _RIGHT
        if top and item.answer and top_count > counts[item.answer - 1]:
            for number in tops:
                sheet.cell(line, 5 + number).fill = _HOT
        sheet.cell(line, 3).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(line, 4).alignment = _CENTER
    _fit_columns(sheet, fixed={"C": 50})
    sheet.freeze_panes = "D2"


def _students_sheet(sheet: Any, graded: Graded) -> None:
    _header(sheet, ["순번", "학번", "이름", "점수", "만점", *TRAP_NAMES, "복습 우선순위"])
    for report in graded.reports:
        traps = dict(report.traps)
        sheet.append(
            [
                report.student.serial,
                None,
                None,
                report.score,
                report.maximum,
                *(traps.get(name, 0) or None for name in TRAP_NAMES),
                " / ".join(report.review),
            ]
        )
        line = sheet.max_row
        _text(sheet.cell(line, 2), report.student.student_id)
        _text(sheet.cell(line, 3), report.student.name)
    _fit_columns(sheet, cap=60)
    sheet.freeze_panes = "D2"


def _totals_sheet(sheet: Any, graded: Graded) -> None:
    _header(sheet, ["함정 유형", "선택 횟수(오답)"])
    for name, count in graded.summary.trap_counts:
        sheet.append([name, count])
    _fit_columns(sheet)


def _excluded_sheet(sheet: Any, excluded: tuple[tuple[str, ...], ...]) -> None:
    _header(sheet, ["시각", "학번(입력값)", "이름", "이유"])
    for entry in excluded:
        sheet.append([None] * 4)
        for column, value in enumerate(entry[:4], 1):
            _text(sheet.cell(sheet.max_row, column), value)
    _fit_columns(sheet)


def _bank_sheet(sheet: Any, bank: QuizBank) -> None:
    source = load_workbook(BytesIO(bank_workbook_bytes(bank)))["문항표"]
    for values in source.iter_rows(values_only=True):
        sheet.append(list(values))
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    _fit_columns(sheet)


def _legend_sheet(sheet: Any) -> None:
    _header(sheet, ["색", "뜻"])
    rows = (
        (_WRONG, "분홍 = 채점결과에서 틀린 답"),
        (_BLANK, "회색 = 채점결과에서 답하지 않음"),
        (_CHECK, "노랑 = 학번이 8자리 숫자가 아님(비고에 입력값)"),
        (_RIGHT, "초록 = 문항 분석의 정답 보기"),
        (_HOT, "빨강 = 문항 분석에서 정답보다 많이 고른 오답"),
    )
    for fill, meaning in rows:
        sheet.append([None, meaning])
        sheet.cell(sheet.max_row, 1).fill = fill
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 50


def result_workbook_bytes(graded: Graded, title: str) -> bytes:
    workbook = Workbook()
    first = workbook.active
    if first is None:
        raise RuntimeError("new workbook must have an active worksheet")
    _result_sheet(first, graded, title)
    _analysis_sheet(workbook.create_sheet(), graded)
    _students_sheet(workbook.create_sheet("학생별 함정"), graded)
    _totals_sheet(workbook.create_sheet("반 전체 함정"), graded)
    if graded.excluded:
        _excluded_sheet(workbook.create_sheet("제외된 응답"), graded.excluded)
    _bank_sheet(workbook.create_sheet("문항표"), graded.bank)
    _legend_sheet(workbook.create_sheet("색 설명"))
    _no_formulas(workbook)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = ["RESULT_FILENAME", "SHEET_NAMES", "result_workbook_bytes"]
