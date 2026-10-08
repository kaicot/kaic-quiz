"""The question bank (문항표): what AI fills in and the feedback reports read.

One row per question: the question and its five options exactly as the form shows them, the
answer, an explanation, why each wrong option is wrong, its trap type from a fixed list, and a
review point. The reader is strict and reports every problem in Korean so the list can be handed
back to the AI as a correction request.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException

from quiz_reporter.errors import Err, ErrorInfo, Ok, Result

BANK_SHEET = "문항표"
GUIDE_SHEET = "설명"
OPTION_COUNT = 5
MAX_QUESTIONS = 100
MAX_PROBLEMS = 40
# Students read 오답이유 as written; longer than this is hard to take in at a glance.
READABLE_LIMIT = 80
TRAP_TYPES: tuple[tuple[str, str], ...] = (
    ("개념 혼동", "비슷한 두 개념을 바꿔 앎 (예: 능동운반/수동운반)"),
    ("용어 혼동", "이름이 비슷한 구조·용어를 바꿈 (예: 희소돌기아교세포/슈반세포)"),
    ("방향·순서 혼동", "이동 방향이나 단계 순서를 반대로 앎 (예: 유입/유출)"),
    ("수치 혼동", "숫자·비율·값을 틀리게 기억 (예: 2/3와 1/3)"),
    ("부분만 맞음", "일부는 맞고 일부가 틀린 보기에 끌림"),
    ("반대 개념", "증가/감소, 촉진/억제처럼 정반대로 앎"),
    ("지나친 일반화", "'모든', '항상'처럼 예외를 무시함"),
    ("사례 적용 오류", "개념은 알지만 임상·상황 문제에 잘못 적용함"),
)
TRAP_NAMES = tuple(name for name, _ in TRAP_TYPES)
BANK_HEADERS: tuple[str, ...] = (
    "번호",
    "단원",
    "문제",
    *(f"보기{number}" for number in range(1, OPTION_COUNT + 1)),
    "정답",
    "해설",
    *(f"오답이유{number}" for number in range(1, OPTION_COUNT + 1)),
    *(f"함정유형{number}" for number in range(1, OPTION_COUNT + 1)),
    "복습포인트",
)
_SPACES = re.compile(r"\s+")


def normalize(text: str) -> str:
    """Compare form and bank text the same way: NFC, single spaces, no outer spaces."""
    return _SPACES.sub(" ", unicodedata.normalize("NFC", text)).strip()


@dataclass(frozen=True, slots=True)
class QuizItem:
    number: int
    unit: str
    question: str
    options: tuple[str, ...]
    answer: int | None
    explanation: str
    reasons: tuple[str, ...]
    traps: tuple[str, ...]
    review: str

    @property
    def has_feedback(self) -> bool:
        wrong = [index for index in range(OPTION_COUNT) if index + 1 != self.answer]
        return bool(self.explanation) and all(self.reasons[index] for index in wrong)


@dataclass(frozen=True, slots=True)
class QuizBank:
    items: tuple[QuizItem, ...]

    @property
    def complete(self) -> bool:
        """Every question has an answer and the feedback a report needs."""
        return all(item.answer is not None and item.has_feedback for item in self.items)


def _problem(text: str) -> ErrorInfo:
    return ErrorInfo("QUIZ_BANK_INVALID", "error.quiz_bank_invalid", None, context={"reason": text})


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return normalize(str(value))


def _number(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    return None


def _answer(value: object) -> int | None:
    """1~5, also written as ①~⑤ or '3번'."""
    if isinstance(value, str):
        text = value.strip().replace("번", "")
        circled = "①②③④⑤"
        if len(text) == 1 and text in circled:
            return circled.index(text) + 1
        value = text
    number = _number(value)
    return number if number is not None and 1 <= number <= OPTION_COUNT else None


def read_bank(rows: list[tuple[object, ...]], *, require_feedback: bool) -> Result[QuizBank]:
    """Rows of the 문항표 sheet, header first. Every problem is listed, numbered by question."""
    if not rows:
        return Err((_problem(f"'{BANK_SHEET}' 시트가 비어 있습니다."),))
    header = [_text(cell) for cell in rows[0]]
    missing = [name for name in BANK_HEADERS if name not in header]
    if missing:
        return Err(
            (
                _problem(
                    f"'{BANK_SHEET}' 시트 첫 줄에 다음 열이 없습니다: {', '.join(missing)}. "
                    f"첫 줄은 {', '.join(BANK_HEADERS)} 이어야 합니다."
                ),
            )
        )
    column = {name: header.index(name) for name in BANK_HEADERS}
    problems: list[str] = []
    items: list[QuizItem] = []
    seen: set[int] = set()
    for row in rows[1:]:
        cells = tuple(row) + (None,) * (len(header) - len(row))
        if all(cell is None or _text(cell) == "" for cell in cells):
            continue

        def get(name: str, cells: tuple[object, ...] = cells) -> object:
            return cells[column[name]]

        number = _number(get("번호"))
        label = f"{number}번 문항" if number is not None else "번호 없는 줄"
        if number is None or not 1 <= number <= MAX_QUESTIONS:
            problems.append(f"{label}: 번호는 1~{MAX_QUESTIONS} 사이의 숫자여야 합니다.")
            continue
        if number in seen:
            problems.append(f"{label}: 번호가 두 번 나옵니다.")
            continue
        seen.add(number)
        question = _text(get("문제"))
        options = tuple(_text(get(f"보기{index}")) for index in range(1, OPTION_COUNT + 1))
        answer = _answer(get("정답"))
        explanation = _text(get("해설"))
        reasons = tuple(_text(get(f"오답이유{index}")) for index in range(1, OPTION_COUNT + 1))
        traps = tuple(_text(get(f"함정유형{index}")) for index in range(1, OPTION_COUNT + 1))
        unit = _text(get("단원"))
        review = _text(get("복습포인트"))
        if not question:
            problems.append(f"{label}: 문제가 비어 있습니다.")
        empty = [str(index + 1) for index, option in enumerate(options) if not option]
        # A bank made from responses alone only knows the options someone chose.
        if empty and require_feedback:
            problems.append(f"{label}: 보기{', 보기'.join(empty)}가 비어 있습니다 (보기는 5개).")
        if len({option for option in options if option}) != len([o for o in options if o]):
            problems.append(f"{label}: 같은 글자의 보기가 있습니다. 보기는 서로 달라야 합니다.")
        if get("정답") not in (None, "") and answer is None:
            problems.append(f"{label}: 정답은 1~5 중 하나의 숫자여야 합니다.")
        if require_feedback:
            if answer is None:
                problems.append(f"{label}: 정답이 비어 있습니다.")
            if not unit:
                problems.append(f"{label}: 단원이 비어 있습니다.")
            if not explanation:
                problems.append(f"{label}: 해설이 비어 있습니다.")
            if not review:
                problems.append(f"{label}: 복습포인트가 비어 있습니다.")
            for index in range(OPTION_COUNT):
                if answer is not None and index + 1 == answer:
                    continue
                if not reasons[index]:
                    problems.append(f"{label}: 오답이유{index + 1}이 비어 있습니다.")
                if traps[index] not in TRAP_NAMES:
                    problems.append(
                        f"{label}: 함정유형{index + 1} '{traps[index]}'은(는) 목록에 없습니다"
                        f" ({', '.join(TRAP_NAMES)} 중 하나)."
                    )
        else:
            for index, trap in enumerate(traps):
                if trap and trap not in TRAP_NAMES:
                    problems.append(
                        f"{label}: 함정유형{index + 1} '{trap}'은(는) 목록에 없습니다"
                        f" ({', '.join(TRAP_NAMES)} 중 하나)."
                    )
        items.append(
            QuizItem(number, unit, question, options, answer, explanation, reasons, traps, review)
        )
    if not items and not problems:
        problems.append(f"'{BANK_SHEET}' 시트에 문항이 없습니다.")
    numbers = sorted(item.number for item in items)
    if items and numbers != list(range(1, len(numbers) + 1)):
        problems.append("번호는 1번부터 빠짐없이 이어져야 합니다.")
    if problems:
        shown = problems[:MAX_PROBLEMS]
        if len(problems) > MAX_PROBLEMS:
            shown.append(f"그 밖에 {len(problems) - MAX_PROBLEMS}건이 더 있습니다.")
        return Err(tuple(_problem(text) for text in shown))
    return Ok(QuizBank(tuple(sorted(items, key=lambda item: item.number))))


def readability_notes(bank: QuizBank) -> list[str]:
    """Feedback text a student may struggle with: too long to take in at a glance."""
    notes: list[str] = []
    for item in bank.items:
        for index, reason in enumerate(item.reasons, 1):
            if len(reason) > READABLE_LIMIT:
                notes.append(
                    f"{item.number}번 문항: 오답이유{index}가 {len(reason)}자로 깁니다."
                    f" 쉬운 말로 {READABLE_LIMIT}자 이내로 줄여 주세요."
                )
        if len(item.explanation) > READABLE_LIMIT * 2:
            notes.append(
                f"{item.number}번 문항: 해설이 {len(item.explanation)}자로 깁니다. 쉬운 말로 줄여 주세요."
            )
    return notes


def parse_bank_bytes(data: bytes, *, require_feedback: bool = True) -> Result[QuizBank]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, KeyError):
        return Err((_problem("엑셀 파일을 열 수 없습니다. .xlsx 파일인지 확인하세요."),))
    try:
        if BANK_SHEET not in workbook.sheetnames:
            return Err((_problem(f"'{BANK_SHEET}' 시트가 없습니다."),))
        sheet: Any = workbook[BANK_SHEET]
        if (sheet.max_row or 0) > MAX_QUESTIONS + 50:
            return Err((_problem(f"'{BANK_SHEET}' 시트가 너무 큽니다."),))
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return read_bank(rows, require_feedback=require_feedback)


def parse_bank(path: str, *, require_feedback: bool = True) -> Result[QuizBank]:
    try:
        data = Path(path).read_bytes()
    except OSError:
        return Err((_problem("문항표 파일을 읽을 수 없습니다."),))
    return parse_bank_bytes(data, require_feedback=require_feedback)


def problems_text(errors: tuple[ErrorInfo, ...]) -> str:
    return "\n".join(f"- {error.context.get('reason', error.code)}" for error in errors)


_HEADER_FILL = PatternFill(fill_type="solid", start_color="FFD9E1F2", end_color="FFD9E1F2")
_TODO_FILL = PatternFill(fill_type="solid", start_color="FFFFEB9C", end_color="FFFFEB9C")


def _append(sheet: Any, values: list[object]) -> None:
    """Append a row keeping text exactly; text that looks like a formula is stored as text."""
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _bank_sheet(workbook: Workbook, bank: QuizBank, title: str = BANK_SHEET) -> Any:
    sheet = workbook.create_sheet(title)
    _append(sheet, list(BANK_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for item in bank.items:
        _append(
            sheet,
            [
                item.number,
                item.unit,
                item.question,
                *item.options,
                item.answer,
                item.explanation,
                *item.reasons,
                *item.traps,
                item.review,
            ],
        )
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.value in (None, ""):
                cell.fill = _TODO_FILL
    widths = {"A": 6, "B": 14, "C": 40}
    for letter, width in widths.items():
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "D2"
    return sheet


def bank_workbook_bytes(bank: QuizBank) -> bytes:
    """A 문항표 workbook; empty cells are yellow so it is clear what the AI still has to fill."""
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    _bank_sheet(workbook, bank)
    guide = workbook.create_sheet(GUIDE_SHEET)
    lines = [
        "문항표: 한 줄에 한 문항. 문제와 보기는 구글 폼에 보이는 글자 그대로 적습니다.",
        "정답: 1~5. 오답이유·함정유형은 정답이 아닌 보기 4개에만 적습니다(정답 칸은 비움).",
        "함정유형은 아래 목록 중 하나만 씁니다.",
        *(f"  {name}: {meaning}" for name, meaning in TRAP_TYPES),
        "노란 칸은 아직 비어 있는 칸입니다.",
    ]
    for line in lines:
        _append(guide, [line])
    guide.column_dimensions["A"].width = 100
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = [
    "BANK_HEADERS",
    "BANK_SHEET",
    "OPTION_COUNT",
    "TRAP_NAMES",
    "TRAP_TYPES",
    "QuizBank",
    "QuizItem",
    "READABLE_LIMIT",
    "bank_workbook_bytes",
    "normalize",
    "parse_bank",
    "parse_bank_bytes",
    "problems_text",
    "read_bank",
    "readability_notes",
]
