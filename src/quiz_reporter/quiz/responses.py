"""Read Google Forms responses as downloaded: the form's CSV or its linked spreadsheet (.xlsx).

The CSV of a quiz-mode form has, per question, the chosen option text, a ``[점수]`` column with
"1.00 / 1" or "0.00 / 1" and an empty ``[의견]`` column. The spreadsheet keeps only the chosen
text. Option text may contain commas, so a cell is one answer and is never split.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from quiz_reporter.errors import Err, ErrorInfo, Ok, Result
from quiz_reporter.quiz.bank import normalize

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 5_000
_TIME_HEADERS = ("타임스탬프", "Timestamp")
_ID_HEADERS = ("학번", "학번(8자리)", "Student ID")
_NAME_HEADERS = ("이름", "성명", "Name")
_META_HEADERS = frozenset({"점수", "총점", "이메일 주소", "Email Address", "Score"})
_SCORE_SUFFIX = "[점수]"
_COMMENT_SUFFIX = "[의견]"
_CSV_TIME = re.compile(
    r"(\d{4})/(\d{1,2})/(\d{1,2})\s+(\d{1,2}):(\d{2}):(\d{2})\s*(오전|오후|AM|PM)?"
    r"(?:\s*GMT([+-]\d{1,2})(?::?(\d{2}))?)?"
)
_ITEM_SCORE = re.compile(r"\s*([0-9.]+)\s*/\s*([0-9.]+)\s*")
KST = timezone(timedelta(hours=9))


@dataclass(frozen=True, slots=True)
class FormResponse:
    row: int
    submitted_at: datetime | None
    student_id: str
    name: str
    answers: tuple[str, ...]
    # Per question: True/False from the CSV's [점수] columns, None when the file has none.
    correct: tuple[bool | None, ...]


@dataclass(frozen=True, slots=True)
class FormResponses:
    source_name: str
    questions: tuple[str, ...]
    rows: tuple[FormResponse, ...]

    @property
    def has_item_scores(self) -> bool:
        return any(value is not None for row in self.rows for value in row.correct)


def _error(reason: str) -> Err:
    return Err(
        (
            ErrorInfo(
                "QUIZ_RESPONSES_INVALID",
                "error.quiz_responses_invalid",
                None,
                context={"reason": reason},
            ),
        )
    )


def parse_csv_time(text: str) -> datetime | None:
    """'2026/01/02 3:04:05 오후 GMT+9' as an aware time (Korea when no zone is given)."""
    match = _CSV_TIME.search(text)
    if match is None:
        return None
    year, month, day, hour, minute, second = (int(match[index]) for index in range(1, 7))
    meridiem = match[7]
    if meridiem in ("오후", "PM") and hour != 12:
        hour += 12
    if meridiem in ("오전", "AM") and hour == 12:
        hour = 0
    zone = KST
    if match[8] is not None:
        offset = int(match[8])
        minutes = int(match[9]) if match[9] else 0
        sign = -1 if offset < 0 else 1
        zone = timezone(timedelta(hours=offset, minutes=sign * minutes))
    try:
        return datetime(year, month, day, hour, minute, second, tzinfo=zone)
    except ValueError:
        return None


def student_id_text(value: object) -> str:
    """Spreadsheets turn IDs into numbers (20260001.0); bring back the digits."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    text = normalize(str(value))
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize(str(value))


def _find(header: list[str], names: tuple[str, ...]) -> int | None:
    for name in names:
        if name in header:
            return header.index(name)
    return None


def read_table(
    source_name: str, header_raw: list[object], body: list[list[object]]
) -> Result[FormResponses]:
    header = [_cell_text(cell) for cell in header_raw]
    id_column = _find(header, _ID_HEADERS)
    name_column = _find(header, _NAME_HEADERS)
    time_column = _find(header, _TIME_HEADERS)
    if id_column is None:
        return _error(
            "응답 파일에서 '학번' 질문을 찾지 못했습니다. 구글 폼에 제목이 '학번'인 질문이 있어야"
            " 합니다."
        )
    skip = {id_column, name_column, time_column} - {None}
    question_columns: list[int] = []
    score_columns: list[int | None] = []
    for index, title in enumerate(header):
        if index in skip or not title or title in _META_HEADERS:
            continue
        if title.endswith((_SCORE_SUFFIX, _COMMENT_SUFFIX)):
            continue
        question_columns.append(index)
        score = f"{title}{_SCORE_SUFFIX}"
        score_columns.append(header.index(score) if score in header else None)
    if not question_columns:
        return _error("응답 파일에서 문항 열을 찾지 못했습니다.")
    rows: list[FormResponse] = []
    for number, raw in enumerate(body, 2):
        cells = list(raw) + [None] * (len(header) - len(raw))
        if all(_cell_text(cell) == "" for cell in cells):
            continue
        when: datetime | None = None
        if time_column is not None:
            stamp = cells[time_column]
            if isinstance(stamp, datetime):
                when = stamp if stamp.tzinfo else stamp.replace(tzinfo=KST)
            elif stamp is not None:
                when = parse_csv_time(str(stamp))
        correct: list[bool | None] = []
        for column in score_columns:
            value = None
            if column is not None:
                match = _ITEM_SCORE.fullmatch(_cell_text(cells[column]))
                if match is not None:
                    value = float(match[1]) > 0
            correct.append(value)
        rows.append(
            FormResponse(
                number,
                when,
                student_id_text(cells[id_column]),
                _cell_text(cells[name_column]) if name_column is not None else "",
                tuple(_cell_text(cells[column]) for column in question_columns),
                tuple(correct),
            )
        )
    if not rows:
        return _error("응답 파일에 응답이 없습니다.")
    return Ok(
        FormResponses(source_name, tuple(header[index] for index in question_columns), tuple(rows))
    )


def read_form_responses(path: str) -> Result[FormResponses]:
    source = Path(path)
    try:
        if source.stat().st_size > MAX_BYTES:
            return _error("응답 파일이 너무 큽니다.")
        data = source.read_bytes()
    except OSError:
        return _error("응답 파일을 읽을 수 없습니다.")
    suffix = source.suffix.lower()
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("cp949")
            except UnicodeDecodeError:
                return _error(
                    "CSV 글자 인코딩을 읽을 수 없습니다. 구글 폼에서 받은 CSV인지 확인하세요."
                )
        table = list(csv.reader(io.StringIO(text)))
        if len(table) > MAX_ROWS:
            return _error("응답이 너무 많습니다.")
        if not table:
            return _error("응답 파일이 비어 있습니다.")
        return read_table(source.name, list(table[0]), [list(row) for row in table[1:]])
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError):
            return _error("엑셀 파일을 열 수 없습니다.")
        try:
            sheet: Any = workbook.worksheets[0]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        if len(rows) > MAX_ROWS:
            return _error("응답이 너무 많습니다.")
        if not rows:
            return _error("응답 파일이 비어 있습니다.")
        return read_table(source.name, rows[0], rows[1:])
    return _error("구글 폼 응답 CSV(.csv) 또는 스프레드시트(.xlsx)를 고르세요.")


__all__ = [
    "KST",
    "FormResponse",
    "FormResponses",
    "parse_csv_time",
    "read_form_responses",
    "read_table",
    "student_id_text",
]
