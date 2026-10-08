"""The teacher's quiz analysis workbook: which traps caught the class, per question and per student."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from quiz_reporter.quiz.bank import TRAP_NAMES, QuizBank, bank_workbook_bytes
from quiz_reporter.quiz.report import CIRCLED, ClassSummary, StudentReport

_BOLD = Font(bold=True)
_HEAD = PatternFill(fill_type="solid", start_color="FFD9E1F2", end_color="FFD9E1F2")
_RIGHT = PatternFill(fill_type="solid", start_color="FFC6EFCE", end_color="FFC6EFCE")
_HOT = PatternFill(fill_type="solid", start_color="FFFFC7CE", end_color="FFFFC7CE")


def _text(cell: object, value: str) -> None:
    cell.value = value  # type: ignore[attr-defined]
    cell.data_type = "s"  # type: ignore[attr-defined]


def _header(sheet: object, values: list[str]) -> None:
    sheet.append(values)  # type: ignore[attr-defined]
    for cell in sheet[1]:  # type: ignore[index]
        cell.font = _BOLD
        cell.fill = _HEAD


def analysis_workbook_bytes(
    bank: QuizBank,
    reports: tuple[StudentReport, ...],
    summary: ClassSummary,
    excluded: tuple[tuple[str, ...], ...],
) -> bytes:
    workbook = Workbook()
    items = workbook.active
    if items is None:
        raise RuntimeError("new workbook must have an active worksheet")
    items.title = "문항별 함정"
    _header(
        items,
        [
            "번호",
            "단원",
            "문제",
            "정답",
            "정답률",
            *(f"보기{CIRCLED[i]} 선택" for i in range(5)),
            "가장 많이 고른 오답",
            "그 함정 유형",
        ],
    )
    for item in bank.items:
        index = item.number - 1
        counts = summary.option_counts[index]
        wrong = [(count, number) for number, count in enumerate(counts, 1) if number != item.answer]
        top_count, top = max(wrong) if wrong else (0, 0)
        line_values: list[object] = [
            item.number,
            item.unit,
            item.question,
            CIRCLED[item.answer - 1] if item.answer else "",
            round(summary.correct_rate[index], 3) if item.answer else None,
            *counts,
            f"{CIRCLED[top - 1]} ({top_count}명)" if top and top_count else "",
            item.traps[top - 1] if top and top_count else "",
        ]
        items.append(line_values)
        line = items.max_row
        _text(items.cell(line, 3), item.question)
        items.cell(line, 5).number_format = "0%"
        if item.answer:
            items.cell(line, 5 + item.answer).fill = _RIGHT
        if top and top_count and item.answer and top_count > counts[item.answer - 1]:
            items.cell(line, 5 + top).fill = _HOT
        items.cell(line, 3).alignment = Alignment(wrap_text=True, vertical="top")
    items.column_dimensions["C"].width = 50
    items.freeze_panes = "D2"

    students = workbook.create_sheet("학생별 함정")
    _header(students, ["순번", "학번", "이름", "점수", "만점", *TRAP_NAMES, "복습 우선순위"])
    for report in reports:
        traps = dict(report.traps)
        students.append(
            [
                report.student.serial,
                None,
                None,
                report.score,
                report.maximum,
                *(traps.get(name, 0) or None for name in TRAP_NAMES),
                " / ".join(report.review),
            ]
        )
        line = students.max_row
        _text(students.cell(line, 2), report.student.student_id)
        _text(students.cell(line, 3), report.student.name)
    students.freeze_panes = "D2"

    totals = workbook.create_sheet("반 전체 함정")
    _header(totals, ["함정 유형", "선택 횟수(오답)"])
    for name, count in summary.trap_counts:
        totals.append([name, count])
    totals.column_dimensions["A"].width = 16

    if excluded:
        left_out = workbook.create_sheet("제외된 응답")
        _header(left_out, ["시각", "학번(입력값)", "이름", "이유"])
        for entry in excluded:
            left_out.append([None] * 4)
            for column, value in enumerate(entry[:4], 1):
                _text(left_out.cell(left_out.max_row, column), value)

    bank_book = load_workbook(BytesIO(bank_workbook_bytes(bank)))
    source = bank_book["문항표"]
    copy = workbook.create_sheet("문항표")
    for values in source.iter_rows(values_only=True):
        copy.append(list(values))
        for cell in copy[copy.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


__all__ = ["analysis_workbook_bytes"]
